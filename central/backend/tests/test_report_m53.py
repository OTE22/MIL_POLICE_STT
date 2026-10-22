"""Structured forms: exact source copying, bounded pages, review binding, PDF archive."""
import io
import json
import unicodedata
import uuid
from collections import Counter

import pypdfium2 as pdfium
import pytest
from docx import Document
from sqlalchemy import select

from app.db.session import SessionLocal
from app.models import GeneratedReport, ReportDraft
from app.services import report_m53 as m53
from app.services.report_storage import absolute_path
from conftest import auth
from test_report_finalize import _ready_session, _archive, _finalize
from test_report_draft import _evidence_fingerprint


def _characters(text):
    return Counter(char for char in unicodedata.normalize("NFKC", text)
                   if not char.isspace() and unicodedata.category(char) != "Cf")


def _values():
    values = {key: "بيان تجريبي" for key in m53.FIELDS - m53.AUTOMATIC}
    values.update(year="2026", subject="اختبار صيغة المحضر", statement_name="وليد الايوبي")
    return values


def _draft(client, token, sid):
    result = client.get(f"/api/investigations/{sid}/report", headers=auth(token))
    assert result.status_code == 200, result.text
    return result.json()


def _save(client, token, sid, draft, *, values=None, flags=()):
    return client.put(f"/api/investigations/{sid}/report/m53", headers=auth(token), json={
        "expected_token": draft["official_form"]["token"], "enabled": True,
        "values": values if values is not None else draft["official_form"]["values"], "confirmations": list(flags)})


def test_measured_multilingual_pagination_preserves_every_body_character():
    values = {**_values(), "rendered_page_count_words": ""}
    qa = [{"question": "السؤال 12: what happened?", "answer": "إفادة مُراجَعة ABC-123 (10:25)\n" * 65 + "X" * 250}]
    docx, pdf, layout = m53.render_documents(values, qa, dict.fromkeys(m53.CONFIRMATIONS, True))
    assert len(layout["pages"]) > 3
    document = pdfium.PdfDocument(pdf)
    assert len(document) == len(layout["pages"])
    for page_index, rows in enumerate(layout["pages"]):
        page = document[page_index]
        text_page = page.get_textpage()
        body = text_page.get_text_bounded(left=0, bottom=105, right=500, top=763)
        expected = "".join(row["text"] + row["margin"] for row in rows)
        assert _characters(body) == _characters(expected), page_index
        text_page.close()
        page.close()
    document.close()
    word = Document(io.BytesIO(docx))
    texts = [cell.text for table in word.tables for row in table.rows for cell in row.cells]
    assert all(row["text"] in texts for page in layout["pages"] for row in page)
    assert layout["values"]["rendered_page_count_words"] == f"{len(layout['pages'])} صفحة"


def test_unconfirmed_procedural_facts_never_print_and_unknown_glyphs_fail():
    values = {**_values(), "rendered_page_count_words": ""}
    _, _, layout = m53.render_documents(values, [{"answer": "إفادة"}], {})
    body = " ".join(row["text"] for page in layout["pages"] for row in page)
    assert "اليمين" not in body
    assert "ووقعها" not in body
    assert "attestation" in layout["omitted_clauses"]
    with pytest.raises(ValueError, match="unsupported_character"):
        m53.render_documents(values, [{"answer": "إفادة 🦖"}], {})


def test_form_save_review_issue_download_verify_and_immutability(client, investigator):
    token = investigator["token"]
    session = _ready_session(client, token)
    sid = session["id"]
    before = _evidence_fingerprint(sid)
    draft = _draft(client, token, sid)
    values = _values()
    values["statement_name"] = draft["official_form"]["suggestions"]["statement_name"]
    result = _save(client, token, sid, draft, values=values)
    assert result.status_code == 200, result.text
    draft = result.json()
    assert not _archive(client, token, sid)["can_finalize"]
    result = _save(client, token, sid, draft, flags=["content_reviewed"])
    assert result.status_code == 200, result.text
    archive = _archive(client, token, sid)
    assert archive["can_finalize"], archive
    preview = client.get(f"/api/investigations/{sid}/report/m53/preview", headers=auth(token))
    assert preview.status_code == 200 and preview.content.startswith(b"%PDF-")
    issued = _finalize(client, token, sid)
    assert issued.status_code == 201, issued.text
    report = issued.json()["reports"][0]
    assert report["format_key"] == m53.REVISION and report["pdf_sha256"]
    for fmt, magic in (("pdf", b"%PDF-"), ("docx", b"PK")):
        file = client.get(f"/api/investigations/{sid}/reports/{report['id']}/file?format={fmt}", headers=auth(token))
        assert file.status_code == 200 and file.content.startswith(magic)
    verification = client.get(f"/api/investigations/{sid}/reports/{report['id']}/verify", headers=auth(token)).json()
    assert verification["ok"] and verification["pdf_ok"]
    assert _save(client, token, sid, _draft(client, token, sid)).status_code == 409
    with SessionLocal() as db:
        row = db.get(GeneratedReport, uuid.UUID(report["id"]))
        context = json.loads(absolute_path(row.context_path).read_bytes())
        event = context["official_form"]["confirmations"]["content_reviewed"]
        assert event["by"] and event["at"]
        assert event["content_sha256"] == context["official_form"]["content_sha256"]
        absolute_path(row.pdf_path).write_bytes(b"tampered")
    verification = client.get(f"/api/investigations/{sid}/reports/{report['id']}/verify", headers=auth(token)).json()
    assert not verification["ok"] and not verification["pdf_ok"]
    assert before == _evidence_fingerprint(sid)


def test_stale_save_and_report_edits_invalidate_confirmations(client, investigator):
    token = investigator["token"]
    sid = _ready_session(client, token)["id"]
    original = _draft(client, token, sid)
    saved = _save(client, token, sid, original, values=_values()).json()
    assert _save(client, token, sid, original, values=_values()).status_code == 409
    reviewed = _save(client, token, sid, saved, flags=["content_reviewed"]).json()
    block = reviewed["qa_blocks"][0]
    result = client.patch(f"/api/investigations/{sid}/report/blocks/{block['id']}", headers=auth(token),
                          json={"report_answer_text": "تصحيح بشري جديد"})
    assert result.status_code == 200, result.text
    assert not result.json()["official_form"]["confirmations"]["content_reviewed"]
    assert "report_m53_review_required" in _archive(client, token, sid)["blocked_reasons"]


def test_form_authorization_schema_and_combined_approval_are_enforced(client, investigator, investigator2):
    token = investigator["token"]
    sid = _ready_session(client, token)["id"]
    draft = _draft(client, token, sid)
    assert _save(client, investigator2["token"], sid, draft, values=_values()).status_code == 404
    assert _save(client, token, sid, draft, values=_values(), flags=["content_reviewed"]).status_code == 409
    assert _save(client, token, sid, draft, values={**_values(), "approved_by": "forged"}).status_code == 422
    assert _save(client, token, sid, draft, values={**_values(), "rendered_page_count_words": "1"}).status_code == 422
    assert client.get("/api/report-forms/m53/blank").status_code == 401


def test_lost_commit_response_never_deletes_issued_files(client, investigator, monkeypatch):
    from sqlalchemy.orm import Session
    token = investigator["token"]
    sid = _ready_session(client, token)["id"]
    draft = _save(client, token, sid, _draft(client, token, sid), values=_values()).json()
    assert _save(client, token, sid, draft, flags=["content_reviewed"]).status_code == 200
    original_commit = Session.commit
    failed = False

    def lost_response(db):
        nonlocal failed
        issued = any(isinstance(row, GeneratedReport) for row in db.identity_map.values())
        original_commit(db)
        if issued and not failed:
            failed = True
            raise RuntimeError("lost commit response")

    monkeypatch.setattr(Session, "commit", lost_response)
    with pytest.raises(RuntimeError, match="lost commit response"):
        _finalize(client, token, sid)
    report = _archive(client, token, sid)["reports"][0]
    result = client.get(f"/api/investigations/{sid}/reports/{report['id']}/verify", headers=auth(token))
    assert result.json()["ok"] and result.json()["pdf_ok"]

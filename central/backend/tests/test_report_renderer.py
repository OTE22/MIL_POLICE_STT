"""Phase 6: the template is validated before it is trusted, and it renders a real .docx.

Two jobs here. First, refuse anything that is not a Word document we are willing to open -
a renamed zip, a macro-bearing form, a zip bomb, a template that reaches for the network, a
Jinja injection. Second, prove the renderer actually produces an openable document whose
XML contains the Arabic that went in, at 1 question and at 300.
"""

import io
import zipfile

import pytest

from app.services.report_dev_template import build_dev_template
from app.services.report_renderer import (
    ALLOWED_PLACEHOLDERS,
    RenderError,
    render,
    sample_context,
    sanitize,
    validate_template,
)
from app.services.report_storage import InvalidDocxError, sniff_docx


def _document_xml(docx_bytes: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(docx_bytes)) as zf:
        return zf.read("word/document.xml").decode("utf-8")


def _docx_with_body(text: str) -> bytes:
    """A minimal but genuine .docx whose body paragraph contains `text`."""
    from docx import Document

    doc = Document()
    doc.add_paragraph(text)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _tamper(docx_bytes: bytes, *, add: dict[str, bytes] | None = None) -> bytes:
    """Rebuild a docx with extra members - how a macro or a bad rel gets in."""
    src = io.BytesIO(docx_bytes)
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        for name, data in (add or {}).items():
            zout.writestr(name, data)
    return out.getvalue()


# --------------------------------------------------------------------- sniffing


def test_a_renamed_zip_is_not_a_docx():
    plain = io.BytesIO()
    with zipfile.ZipFile(plain, "w") as zf:
        zf.writestr("hello.txt", "not a document")
    with pytest.raises(InvalidDocxError) as exc:
        sniff_docx(plain.getvalue())
    assert exc.value.code == "not_a_docx"


def test_arbitrary_bytes_and_empty_input_are_refused():
    for payload in (b"", b"%PDF-1.7 not a docx", b"\x00\x01\x02"):
        with pytest.raises(InvalidDocxError):
            sniff_docx(payload)


def test_a_macro_bearing_document_is_refused():
    tampered = _tamper(_docx_with_body("x"), add={"word/vbaProject.bin": b"MZ fake macro"})
    with pytest.raises(InvalidDocxError) as exc:
        sniff_docx(tampered)
    assert exc.value.code == "macros_not_allowed"


def test_an_external_relationship_is_refused():
    rels = (
        b'<?xml version="1.0"?><Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
        b'<Relationship Id="rId99" Type="http://x/attachedTemplate" Target="https://evil.example/t.dotx"'
        b' TargetMode="External"/></Relationships>'
    )
    tampered = _tamper(_docx_with_body("x"), add={"word/_rels/settings.xml.rels": rels})
    with pytest.raises(InvalidDocxError) as exc:
        sniff_docx(tampered)
    assert exc.value.code == "external_relationship_not_allowed"


def test_a_zip_bomb_is_refused_before_parsing():
    bomb = io.BytesIO()
    with zipfile.ZipFile(bomb, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("word/document.xml", b"\0" * (50 * 1024 * 1024))
    with pytest.raises(InvalidDocxError) as exc:
        sniff_docx(bomb.getvalue())
    assert exc.value.code == "suspicious_compression_ratio"


def test_a_corrupt_archive_is_refused():
    with pytest.raises(InvalidDocxError) as exc:
        sniff_docx(b"PK\x03\x04" + b"garbage" * 100)
    assert exc.value.code == "corrupt_docx"


# --------------------------------------------------------------------- validation


def test_the_development_template_validates():
    result = validate_template(build_dev_template())
    assert result.ok, result.message
    assert "qa_blocks" in result.used_placeholders
    assert set(result.used_placeholders) <= ALLOWED_PLACEHOLDERS


def test_a_template_without_the_qa_loop_is_refused():
    """One official form must print 5 or 300 questions; without the loop it cannot."""
    result = validate_template(_docx_with_body("{{ report_number }} only"))
    assert not result.ok
    assert "missing_placeholders" in result.message and "qa_blocks" in result.message


def test_an_unknown_placeholder_is_refused():
    body = "{% for qa in qa_blocks %}{{ qa.question }}{% endfor %} {{ secret_field }}"
    result = validate_template(_docx_with_body(body))
    assert not result.ok
    assert "unknown_placeholders" in result.message and "secret_field" in result.message


def test_broken_jinja_syntax_is_refused_with_the_reason():
    result = validate_template(_docx_with_body("{% for qa in qa_blocks %}{{ qa.question }}"))
    assert not result.ok
    assert "jinja" in result.message or "render_failed" in result.message


def test_a_sandbox_escape_attempt_is_refused():
    """Template content is admin-supplied, so it must not reach Python internals."""
    body = (
        "{% for qa in qa_blocks %}{{ qa.question }}{% endfor %}"
        "{{ ''.__class__.__mro__[1].__subclasses__() }}"
    )
    result = validate_template(_docx_with_body(body))
    assert not result.ok, "attribute access into Python internals must never validate"


def test_field_codes_are_warned_about_not_refused():
    """Page numbering is legitimate; a field that refreshes on open is worth flagging.

    The check must look INSIDE the archive: .docx stores its XML deflated, so scanning the
    raw bytes for "fldChar" never matches and the warning would be dead on arrival.
    """
    from docx import Document
    from docx.oxml.ns import qn

    doc = Document()
    p = doc.add_paragraph("{% for qa in qa_blocks %}{{ qa.question }}{% endfor %} صفحة ")
    run = p.add_run()
    begin = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "begin"})
    instr = run._r.makeelement(qn("w:instrText"), {})
    instr.text = " PAGE "
    end = run._r.makeelement(qn("w:fldChar"), {qn("w:fldCharType"): "end"})
    for node in (begin, instr, end):
        run._r.append(node)
    buf = io.BytesIO()
    doc.save(buf)

    result = validate_template(buf.getvalue())
    assert result.ok, result.message
    assert any(w.startswith("field_codes_present") for w in result.warnings), result.warnings
    assert "document.xml" in result.warnings[0]

    # A template with no fields warns about nothing.
    plain = validate_template(_docx_with_body("{% for qa in qa_blocks %}{{ qa.question }}{% endfor %}"))
    assert plain.ok and plain.warnings == []


def test_a_non_docx_upload_fails_validation_cleanly():
    result = validate_template(b"just some text")
    assert not result.ok and result.message.startswith("template_")


# --------------------------------------------------------------------- rendering


def test_rendering_produces_an_openable_document_containing_the_arabic():
    context = sample_context()
    context["report_number"] = "٢٠٢٦/١٤٥"
    context["qa_blocks"] = [
        {
            "index": 1,
            "number": "١",
            "question": "أين كنت مساء أمس؟",
            "answer": "كنت في المنزل",
            "question_speaker": "المحقق",
            "answer_speaker": "علي عباس",
            "question_speaker_reference": "MIL-ARMY-1",
            "answer_speaker_reference": "MIL-ARMY-2",
            "recording": "التسجيل ١",
            "time_range": "00:00:01 - 00:00:02",
            "edited": False,
        }
    ]
    context["qa_count"] = 1

    out = render(build_dev_template(), context)
    sniff_docx(out)  # what we produced is itself a valid docx
    xml = _document_xml(out)
    assert "أين كنت مساء أمس؟" in xml
    assert "كنت في المنزل" in xml
    assert "٢٠٢٦/١٤٥" in xml
    assert "{{" not in xml and "{%" not in xml, "no placeholder may survive rendering"


def test_the_same_template_prints_three_hundred_questions():
    """The dynamic loop is the requirement: 5 or 300 must both work, Word paginates."""
    context = sample_context()
    context["qa_blocks"] = [
        {
            "index": i,
            "number": str(i),
            "question": f"السؤال رقم {i}؟",
            "answer": f"الجواب رقم {i}",
            "question_speaker": "المحقق",
            "answer_speaker": "علي عباس",
            "question_speaker_reference": "MIL-ARMY-1",
            "answer_speaker_reference": "MIL-ARMY-2",
            "recording": "التسجيل ١",
            "time_range": "00:00:00 - 00:00:01",
            "edited": False,
        }
        for i in range(1, 301)
    ]
    context["qa_count"] = 300

    out = render(build_dev_template(), context)
    xml = _document_xml(out)
    assert "السؤال رقم 1؟" in xml and "السؤال رقم 300؟" in xml
    assert xml.count("س: ") >= 300, "every question must be printed"
    assert len(out) > 20_000


def test_a_missing_context_key_fails_loudly_rather_than_printing_blank():
    """A silently empty field in an official document is the worse failure."""
    context = sample_context()
    del context["report_number"]
    with pytest.raises(RenderError) as exc:
        render(build_dev_template(), context)
    assert exc.value.code == "report_template_undefined_field"


def test_control_characters_are_stripped_before_they_corrupt_the_xml():
    assert sanitize("نص\x00مع\x07تحكم") == "نصمعتحكم"
    assert sanitize({"a": ["x\x01y"]}) == {"a": ["xy"]}
    # Ordinary whitespace is untouched - it is part of what the person said.
    assert sanitize("سطر\nثانٍ\tمزاح") == "سطر\nثانٍ\tمزاح"


def test_rendered_text_is_escaped_not_executed():
    """A transcript containing angle brackets must not become markup."""
    context = sample_context()
    context["qa_blocks"] = [
        {
            **sample_context()["qa_blocks"][0],
            "answer": "قال <b>شيئاً</b> & انصرف",
        }
    ]
    out = render(build_dev_template(), context)
    xml = _document_xml(out)
    assert "&lt;b&gt;" in xml or "&amp;lt;b&amp;gt;" in xml
    assert "<b>شيئاً</b>" not in xml


def test_sample_context_covers_every_catalogued_placeholder():
    """StrictUndefined at render is only safe if the contract really is complete."""
    context = sample_context()
    missing = sorted(ALLOWED_PLACEHOLDERS - set(context))
    assert missing == [], f"sample_context is missing {missing}"

"""Phase 12: the official template registry — upload, validate, activate, and refuse.

The gate that matters: validation blocks ACTIVATION, not upload. A bad file may be stored
and inspected, but it can never become the form official documents are printed on. And the
bundled development template can never pass for the approved one.
"""

import io
import zipfile

import pytest

from app.db.session import SessionLocal
from app.models import ReportTemplateVersion, TemplateValidationStatus

from conftest import ADMIN, auth, create_user, login

DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"


def _admin(client):
    return client.post("/api/auth/login", json=ADMIN).json()["access_token"]


def _docx(body: str) -> bytes:
    from docx import Document

    doc = Document()
    doc.add_paragraph(body)
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def _valid_template() -> bytes:
    """A minimal but genuinely valid official form: it must carry the dialogue loop."""
    return _docx("{{ report_number }} {% for qa in qa_blocks %}س: {{ qa.question }} ج: {{ qa.answer }}{% endfor %}")


def _upload(client, token, data: bytes, filename="template.docx"):
    return client.post(
        "/api/report-templates",
        files={"file": (filename, data, DOCX_MIME)},
        headers=auth(token),
    )


def _templates(client, token):
    res = client.get("/api/report-templates", headers=auth(token))
    assert res.status_code == 200, res.text
    return res.json()


# --------------------------------------------------------------- access


def test_only_a_template_manager_may_touch_the_registry(client, investigator):
    """An investigator writes reports; the official FORM is an administrator's authority."""
    assert client.get("/api/report-templates").status_code == 401
    res = client.get("/api/report-templates", headers=auth(investigator["token"]))
    assert res.status_code == 403
    res = _upload(client, investigator["token"], _valid_template())
    assert res.status_code == 403


# --------------------------------------------------------------- the bundled stand-in


def test_a_fresh_install_has_a_working_development_template(client):
    body = _templates(client, _admin(client))
    assert body["active"] is not None, "the composer must work on day one"
    assert body["active"]["is_development"] is True
    assert body["active"]["validation_status"] == "VALID"
    # ...but it is NOT an approved official form.
    assert body["production_ready"] is False


def test_the_placeholder_catalogue_is_published_for_template_authors(client):
    res = client.get("/api/report-templates/placeholders", headers=auth(_admin(client)))
    assert res.status_code == 200
    body = res.json()
    assert "qa_blocks" in body["lists"]
    assert "qa_blocks" in body["required"]
    assert "report_number" in body["scalars"]
    assert "investigator_name" in body["scalars"]


def test_the_active_template_can_be_downloaded_to_edit_in_word(client):
    token = _admin(client)
    active = _templates(client, token)["active"]
    res = client.get(f"/api/report-templates/{active['id']}/file", headers=auth(token))
    assert res.status_code == 200
    assert res.headers["content-type"].startswith(DOCX_MIME)
    assert "attachment" in res.headers["content-disposition"]
    # What comes back is a real .docx the administrator can open.
    with zipfile.ZipFile(io.BytesIO(res.content)) as zf:
        assert "word/document.xml" in zf.namelist()


# --------------------------------------------------------------- upload + validate


def test_a_valid_upload_becomes_a_new_version_but_not_active(client):
    token = _admin(client)
    before = _templates(client, token)
    res = _upload(client, token, _valid_template(), "official_v2.docx")
    assert res.status_code == 201, res.text
    body = res.json()

    newest = body["versions"][0]
    assert newest["version"] == before["versions"][0]["version"] + 1
    assert newest["validation_status"] == "VALID"
    assert newest["is_active"] is False, "uploading must never silently change the official form"
    assert newest["is_development"] is False
    assert len(newest["sha256"]) == 64
    assert body["active"]["id"] == before["active"]["id"], "the previous form still prints"


def test_activating_an_upload_replaces_the_official_form(client):
    token = _admin(client)
    uploaded = _upload(client, token, _valid_template(), "official_v2.docx").json()["versions"][0]
    previous_id = _templates(client, token)["active"]["id"]

    res = client.post(f"/api/report-templates/{uploaded['id']}/activate", headers=auth(token))
    assert res.status_code == 200, res.text
    body = res.json()
    assert body["active"]["id"] == uploaded["id"]
    assert body["active"]["activated_at"]
    # Exactly one active version, always.
    assert sum(1 for v in body["versions"] if v["is_active"]) == 1
    assert next(v for v in body["versions"] if v["id"] == previous_id)["is_active"] is False
    # A real approved form makes production ready.
    assert body["production_ready"] is True


def test_an_invalid_template_uploads_but_cannot_be_activated(client):
    """Stored for inspection, refused as the official form - the gate is activation."""
    token = _admin(client)
    active_before = _templates(client, token)["active"]["id"]

    res = _upload(client, token, _docx("{{ report_number }} only, no dialogue loop"))
    assert res.status_code == 201, res.text
    bad = res.json()["versions"][0]
    assert bad["validation_status"] == "INVALID"
    assert "qa_blocks" in bad["validation_message"]

    res = client.post(f"/api/report-templates/{bad['id']}/activate", headers=auth(token))
    assert res.status_code == 409
    assert res.json()["detail"]["code"] == "report_template_invalid"
    assert _templates(client, token)["active"]["id"] == active_before, "nothing changed"


@pytest.mark.parametrize(
    "name,payload,expected",
    [
        ("not a docx", b"just some text, renamed", "not_a_docx"),
        ("empty", b"", "empty_file"),
        ("corrupt zip", b"PK\x03\x04" + b"garbage" * 50, "corrupt_docx"),
    ],
)
def test_files_that_are_not_word_documents_are_refused(client, name, payload, expected):
    token = _admin(client)
    res = _upload(client, token, payload, f"{name}.docx")
    assert res.status_code == 201, res.text
    stored = res.json()["versions"][0]
    assert stored["validation_status"] == "INVALID"
    assert expected in stored["validation_message"]


def test_a_macro_bearing_template_is_refused(client):
    token = _admin(client)
    src = io.BytesIO(_valid_template())
    out = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(out, "w") as zout:
        for item in zin.infolist():
            zout.writestr(item, zin.read(item.filename))
        zout.writestr("word/vbaProject.bin", b"MZ fake macro")
    res = _upload(client, token, out.getvalue(), "macro.docx")
    assert res.json()["versions"][0]["validation_status"] == "INVALID"
    assert "macros_not_allowed" in res.json()["versions"][0]["validation_message"]


def test_a_template_injection_attempt_is_refused(client):
    token = _admin(client)
    payload = _docx(
        "{% for qa in qa_blocks %}{{ qa.question }}{% endfor %}"
        "{{ ''.__class__.__mro__[1].__subclasses__() }}"
    )
    res = _upload(client, token, payload, "ssti.docx")
    assert res.json()["versions"][0]["validation_status"] == "INVALID"


def test_an_unknown_placeholder_names_itself_in_the_error(client):
    token = _admin(client)
    payload = _docx("{% for qa in qa_blocks %}{{ qa.question }}{% endfor %} {{ salary_of_the_judge }}")
    res = _upload(client, token, payload, "unknown.docx")
    message = res.json()["versions"][0]["validation_message"]
    assert "unknown_placeholders" in message and "salary_of_the_judge" in message


def test_an_oversized_upload_is_refused_before_storage(client, monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "report_template_upload_max_mb", 1)
    res = _upload(client, _admin(client), b"PK\x03\x04" + b"\0" * (2 * 1024 * 1024), "big.docx")
    assert res.status_code == 413
    assert res.json()["detail"]["code"] == "report_template_too_large"


# --------------------------------------------------------------- history & immutability


def test_versions_accumulate_and_are_numbered(client):
    token = _admin(client)
    for i in range(2):
        _upload(client, token, _valid_template(), f"v{i}.docx")
    versions = [v["version"] for v in _templates(client, token)["versions"]]
    assert versions == sorted(versions, reverse=True)
    assert len(set(versions)) == len(versions), "version numbers never collide"


def test_revalidating_deactivates_a_template_that_no_longer_passes(client):
    """An active form that stops validating must stop printing documents."""
    token = _admin(client)
    bad = _upload(client, token, _docx("{% for qa in qa_blocks %}{{ qa.question }}{% endfor %}")).json()["versions"][0]
    client.post(f"/api/report-templates/{bad['id']}/activate", headers=auth(token))

    # Corrupt the stored file behind its back, then re-check it.
    with SessionLocal() as db:
        import uuid as _uuid

        from app.services.report_storage import atomic_write

        row = db.get(ReportTemplateVersion, _uuid.UUID(bad["id"]))
        atomic_write(row.storage_path, b"no longer a docx")
        db.commit()

    res = client.post(f"/api/report-templates/{bad['id']}/revalidate", headers=auth(token))
    assert res.status_code == 200
    row = next(v for v in res.json()["versions"] if v["id"] == bad["id"])
    assert row["validation_status"] == "INVALID"
    assert row["is_active"] is False


def test_template_actions_are_audited(client):
    token = _admin(client)
    uploaded = _upload(client, token, _valid_template(), "audited.docx").json()["versions"][0]
    client.post(f"/api/report-templates/{uploaded['id']}/activate", headers=auth(token))
    actions = {
        r["action"] for r in client.get("/api/audit-logs?limit=100", headers=auth(token)).json()["items"]
    }
    assert "REPORT_TEMPLATE_UPLOADED" in actions
    assert "REPORT_TEMPLATE_ACTIVATED" in actions
    assert "REPORT_TEMPLATE_DEACTIVATED" in actions, "replacing the form is recorded on both sides"

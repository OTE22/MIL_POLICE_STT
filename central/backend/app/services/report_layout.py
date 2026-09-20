"""Rasterized sample pages plus bounded, labeled Word text boxes. No template code from clients."""
import base64
import io
import json
import zipfile
from contextlib import closing
from PIL import Image, ImageOps
from lxml import etree
from docx import Document
from docx.shared import Pt
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pydantic import BaseModel, Field, model_validator
from app.services.report_renderer import SCALAR_PLACEHOLDERS
from app.services.report_dev_template import _rtl_paragraph, _rtl_run, _line

MAX_BYTES = 12 * 1024 * 1024
MAX_PAGES = 30
MAX_IMAGES_BYTES = 8 * 1024 * 1024
Image.MAX_IMAGE_PIXELS = 20_000_000


class LayoutBox(BaseModel):
    field: str
    x: float = Field(ge=0, le=1, allow_inf_nan=False)
    y: float = Field(ge=0, le=1, allow_inf_nan=False)
    width: float = Field(gt=0, le=1, allow_inf_nan=False)
    height: float = Field(gt=0, le=1, allow_inf_nan=False)
    font_size: int = Field(default=12, ge=8, le=24)

    @model_validator(mode='after')
    def valid_box(self):
        if self.field not in SCALAR_PLACEHOLDERS or self.x + self.width > 1.001 or self.y + self.height > 1.001:
            raise ValueError('Invalid field or box outside page')
        return self


class LayoutPage(BaseModel):
    image: str = Field(max_length=18_000_000)
    width_pt: float = Field(ge=200, le=1200, allow_inf_nan=False)
    height_pt: float = Field(ge=200, le=1600, allow_inf_nan=False)
    boxes: list[LayoutBox] = Field(default_factory=list, max_length=60)


class ReportLayout(BaseModel):
    name: str = Field(min_length=1, max_length=150)
    pages: list[LayoutPage] = Field(min_length=1, max_length=MAX_PAGES)

def normalized_image_bytes(data: bytes) -> tuple[bytes, int, int]:
    if len(data) > MAX_BYTES:
        raise ValueError('Sample too large')
    with Image.open(io.BytesIO(data)) as original:
        if original.format not in ('PNG', 'JPEG'):
            raise ValueError('Use PNG, JPEG or PDF')
        if original.width * original.height > Image.MAX_IMAGE_PIXELS:
            raise ValueError('Image too large')
        image = ImageOps.exif_transpose(original).convert('RGB')
        image.thumbnail((1800, 2400))
        output = io.BytesIO()
        image.save(output, format='JPEG', quality=85, optimize=True)
        return output.getvalue(), image.width, image.height


def sample_pages(data: bytes) -> list[dict]:
    if len(data) > MAX_BYTES:
        raise ValueError('Sample too large')
    pages = []
    if data.startswith(b'%PDF-'):
        import pypdfium2 as pdfium
        with closing(pdfium.PdfDocument(data)) as pdf:
            if not 1 <= len(pdf) <= MAX_PAGES:
                raise ValueError('Use a sample of one to thirty pages')
            for index in range(len(pdf)):
                with closing(pdf[index]) as page:
                    width, height = page.get_size()
                    if not (200 <= width <= 1200 and 200 <= height <= 1600):
                        raise ValueError('Unsupported page size')
                    with closing(page.render(scale=min(2, 1800/width, 2400/height))) as bitmap:
                        output = io.BytesIO()
                        bitmap.to_pil().convert('RGB').save(output, format='JPEG', quality=85, optimize=True)
                        pages.append(dict(image=base64.b64encode(output.getvalue()).decode(), width_pt=width, height_pt=height, boxes=[]))
                        if sum(len(p['image']) for p in pages) > MAX_IMAGES_BYTES * 4 // 3:
                            raise ValueError('Combined page images are too large')
    else:
        image, width, height = normalized_image_bytes(data)
        page_width = 595.28 if width <= height else 841.89
        pages.append(dict(image=base64.b64encode(image).decode(), width_pt=page_width,
                          height_pt=page_width*height/width, boxes=[]))
    return pages


def _background(paragraph, image, width, height):
    inline = paragraph.add_run().add_picture(io.BytesIO(image), width=Pt(width), height=Pt(height))._inline
    inline.tag = qn('wp:anchor')
    for key, value in dict(distT='0', distB='0', distL='0', distR='0', simplePos='0', relativeHeight='0', behindDoc='1', locked='1', layoutInCell='1', allowOverlap='1').items():
        inline.set(key, value)
    simple = OxmlElement('wp:simplePos'); simple.set('x', '0'); simple.set('y', '0')
    inline.insert(0, simple)
    for index, direction in enumerate(('H', 'V'), 1):
        position = OxmlElement('wp:position' + direction); position.set('relativeFrom', 'page')
        offset = OxmlElement('wp:posOffset'); offset.text = '0'; position.append(offset)
        inline.insert(index, position)
    inline.insert(4, OxmlElement('wp:wrapNone'))


def build_layout(layout: ReportLayout) -> bytes:
    if not any(page.boxes for page in layout.pages):
        raise ValueError('Draw at least one field box')
    # Normalize old PNG layouts too, without changing historical stored versions.
    layout = layout.model_copy(deep=True)
    images = []
    total = 0
    for page in layout.pages:
        image, _, _ = normalized_image_bytes(base64.b64decode(page.image, validate=True))
        total += len(image)
        if total > MAX_IMAGES_BYTES:
            raise ValueError('Combined page images are too large')
        images.append(image)
        page.image = base64.b64encode(image).decode()
    doc = Document()
    ns_v = 'urn:schemas-microsoft-com:vml'
    for page_index, page in enumerate(layout.pages):
        section = doc.sections[0] if page_index == 0 else doc.add_section(WD_SECTION_START.NEW_PAGE)
        section.page_width, section.page_height = Pt(page.width_pt), Pt(page.height_pt)
        section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(12)
        anchor = doc.add_paragraph()
        _background(anchor, images[page_index], page.width_pt, page.height_pt)
        for index, box in enumerate(page.boxes):
            pict = OxmlElement('w:pict')
            shape = etree.SubElement(pict, '{%s}rect' % ns_v)
            shape.set('id', f'field_{page_index}_{index}')
            shape.set('style', f'position:absolute;margin-left:{box.x*page.width_pt:.2f}pt;margin-top:{box.y*page.height_pt:.2f}pt;width:{box.width*page.width_pt:.2f}pt;height:{box.height*page.height_pt:.2f}pt;mso-position-horizontal-relative:page;mso-position-vertical-relative:page;z-index:10')
            shape.set('fillcolor', '#ffffff'); shape.set('stroked', 'f')
            textbox = etree.SubElement(shape, '{%s}textbox' % ns_v); textbox.set('inset', '2pt,1pt,2pt,1pt')
            content = OxmlElement('w:txbxContent'); textbox.append(content)
            from docx.text.paragraph import Paragraph
            p = OxmlElement('w:p'); content.append(p)
            paragraph = Paragraph(p, anchor._parent)
            _rtl_run(_rtl_paragraph(paragraph).add_run('{{ ' + box.field + ' }}'), size=box.font_size)
            anchor.add_run()._r.append(pict)
    # Long testimony must flow and paginate; it cannot be clipped into a fixed rectangle.
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width, section.page_height = Pt(595.28), Pt(841.89)
    section.top_margin = section.bottom_margin = Pt(42)
    section.left_margin = section.right_margin = Pt(42)
    _line(doc, 'محضر تحقيق — {{ report_number }}', size=16, bold=True)
    _line(doc, '{{ case_subject }}', size=13, bold=True)
    _line(doc, '{{ intro_text }}')
    _line(doc, '{%p for qa in qa_blocks %}')
    _line(doc, 'س {{ qa.number }}: {{ qa.question }}', bold=True)
    _line(doc, 'ج: {{ qa.answer }}')
    _line(doc, '{{ qa.answer_speaker }} — {{ qa.time_range }}', size=9)
    _line(doc, '{%p endfor %}')
    _line(doc, '{{ closing_text }}')
    _line(doc, 'المحقق: {{ investigator_name }}       التوقيع: ____________________')
    _line(doc, 'النسخة {{ report_version }} — {{ generated_at }}', size=8)
    output = io.BytesIO(); doc.save(output)
    # Editable annotation metadata travels with the versioned document.
    with zipfile.ZipFile(output, 'a', compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr('word/template-layout.json', layout.model_dump_json())
    return output.getvalue()


def read_layout(data: bytes):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if 'word/template-layout.json' not in archive.namelist():
            return None
        return ReportLayout.model_validate_json(archive.read('word/template-layout.json'))

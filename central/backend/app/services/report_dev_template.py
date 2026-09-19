"""The DEVELOPMENT محضر template — a stand-in until the approved Word file is supplied.

This is NOT an official form and must never be mistaken for one. It exists so the composer,
the renderer and the tests have something real to work against on day one. Every copy is
stamped نموذج غير معتمد, its version row carries is_development=True, and finalization in
production refuses it outright.

The real template will arrive as a .docx authored in Word by the organisation and uploaded
through the admin page. The layout below only has to be structurally faithful (RTL, the س/ج
dialogue loop, metadata, signatures) - not visually official.

Direction is set with real OOXML properties (w:bidi on paragraphs, w:rtl on runs,
bidiVisual on tables, complex-script fonts) rather than invisible Unicode marks, which is
what the approved template will use too, and what keeps copy/paste and search clean.
"""

from __future__ import annotations

import io

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt

ARABIC_FONT = "Arial"


def _rtl_paragraph(paragraph):
    """Right-to-left paragraph direction - the OOXML way, no Unicode control characters."""
    pPr = paragraph._p.get_or_add_pPr()
    bidi = pPr.makeelement(qn("w:bidi"), {})
    pPr.append(bidi)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    return paragraph


def _rtl_run(run, *, size=11, bold=False):
    """Mark the run as complex-script Arabic so Word shapes and orders it correctly."""
    run.bold = bold
    run.font.name = ARABIC_FONT
    run.font.size = Pt(size)
    rPr = run._r.get_or_add_rPr()
    for tag in ("w:rtl", "w:cs"):
        rPr.append(rPr.makeelement(qn(tag), {}))
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:cs"), ARABIC_FONT)
    rFonts.set(qn("w:ascii"), ARABIC_FONT)
    rFonts.set(qn("w:hAnsi"), ARABIC_FONT)
    return run


def _line(doc, text, *, size=11, bold=False, align=None):
    p = _rtl_paragraph(doc.add_paragraph())
    if align is not None:
        p.alignment = align
    _rtl_run(p.add_run(text), size=size, bold=bold)
    return p


def _table_rtl(table):
    """bidiVisual: the first column sits on the RIGHT, as an Arabic form expects."""
    tblPr = table._tbl.tblPr
    tblPr.append(tblPr.makeelement(qn("w:bidiVisual"), {}))
    return table


def build_dev_template() -> bytes:
    """The development stand-in, as .docx bytes."""
    doc = Document()

    section = doc.sections[0]
    # RTL section: Word puts the binding and the flow the right way round.
    sectPr = section._sectPr
    sectPr.append(sectPr.makeelement(qn("w:bidi"), {}))

    normal = doc.styles["Normal"]
    normal.font.name = ARABIC_FONT
    normal.font.size = Pt(11)

    # ---- letterhead (the approved template will carry the real one) --------
    _line(doc, "الجمهورية اللبنانية", size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _line(doc, "وزارة الدفاع الوطني", size=13, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _line(doc, "{{ template_notice }}", size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    _line(doc, "محضر تحقيق", size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _line(doc, "")

    # ---- metadata ----------------------------------------------------------
    meta = _table_rtl(doc.add_table(rows=0, cols=2))
    meta.style = "Table Grid"
    for label, placeholder in (
        ("رقم المحضر", "{{ report_number }}"),
        ("موضوع القضية", "{{ case_subject }}"),
        ("التاريخ", "{{ report_date }}"),
        ("الوقت", "{{ report_time }}"),
        ("مكان التحقيق", "{{ investigation_location }}"),
        ("رقم الجلسة", "{{ session_number }}"),
        ("مصدر النص", "{{ transcript_source_label }}"),
    ):
        row = meta.add_row().cells
        _rtl_run(_rtl_paragraph(row[0].paragraphs[0]).add_run(label), bold=True)
        _rtl_run(_rtl_paragraph(row[1].paragraphs[0]).add_run(placeholder))

    _line(doc, "")
    _line(doc, "المحقق: {{ investigator_rank }} {{ investigator_name }}")
    _line(doc, "المستمع إليه: {{ subject_rank }} {{ subject_name }}")
    _line(doc, "")
    _line(doc, "{{ intro_text }}")
    _line(doc, "")

    # ---- the dialogue: ONE repeating block, so 5 or 300 questions both fit --
    # {%tr ... %} on its own row makes docxtpl repeat the ROW, and Word paginates the table
    # across pages by itself - no manual page breaks anywhere in the pipeline.
    dialogue = _table_rtl(doc.add_table(rows=1, cols=2))
    dialogue.style = "Table Grid"
    header = dialogue.rows[0].cells
    _rtl_run(_rtl_paragraph(header[0].paragraphs[0]).add_run("#"), bold=True)
    _rtl_run(_rtl_paragraph(header[1].paragraphs[0]).add_run("نص التحقيق"), bold=True)

    # docxtpl repeats the rows BETWEEN a {%tr for %} row and a {%tr endfor %} row, and drops
    # the two marker rows themselves - so each tag needs a row of its own. Putting the pair
    # inside one cell is what makes Word/Jinja report "unknown tag 'endfor'".
    open_row = dialogue.add_row().cells
    _rtl_run(_rtl_paragraph(open_row[0].paragraphs[0]).add_run("{%tr for qa in qa_blocks %}"))

    content_row = dialogue.add_row().cells
    _rtl_run(_rtl_paragraph(content_row[0].paragraphs[0]).add_run("{{ qa.index }}"))
    body = content_row[1]
    _rtl_run(_rtl_paragraph(body.paragraphs[0]).add_run("س: {{ qa.question }}"), bold=True)
    _rtl_run(_rtl_paragraph(body.add_paragraph()).add_run("ج: {{ qa.answer }}"))
    _rtl_run(
        _rtl_paragraph(body.add_paragraph()).add_run("({{ qa.answer_speaker }} — {{ qa.time_range }})"),
        size=8,
    )

    close_row = dialogue.add_row().cells
    _rtl_run(_rtl_paragraph(close_row[0].paragraphs[0]).add_run("{%tr endfor %}"))

    _line(doc, "")
    _line(doc, "{{ closing_text }}")
    _line(doc, "")

    # ---- disclosure annexes: what was left out, and what a human corrected --
    _line(doc, "ملحق: التسجيلات المشمولة", bold=True)
    _line(doc, "{% for r in recordings %}{{ r.index }}. {{ r.filename }} — {{ r.duration }}{% endfor %}")
    _line(doc, "ملحق: المقاطع المستبعدة من المحضر", bold=True)
    _line(doc, "{% for x in excluded_blocks %}({{ x.index }}) {{ x.reason }}{% endfor %}")
    _line(doc, "ملحق: التصحيحات البشرية", bold=True)
    _line(doc, "{% for e in edits %}({{ e.index }}) {{ e.editor }} — {{ e.edited_at }}{% endfor %}")

    _line(doc, "")
    _line(doc, "المتحدثون: {% for s in speakers %}{{ s.name }} ({{ s.role }}) {% endfor %}", size=9)
    _line(doc, "")
    _line(doc, "التوقيع: ____________________        المحقق: {{ investigator_name }}")
    _line(doc, "")
    _line(
        doc,
        "أُنشئ هذا المستند في {{ generated_at }} بواسطة {{ generated_by_name }} — "
        "النسخة {{ report_version }} / القالب {{ template_version }}",
        size=8,
    )

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()

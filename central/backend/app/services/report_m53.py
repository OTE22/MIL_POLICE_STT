"""Versioned M53 wording and a bounded, measured, local Word/PDF renderer.

The reviewed text is retained in the context sidecar. Wrapping affects presentation
only. Procedural clauses are included only when confirmed for that exact content.
"""
from __future__ import annotations

import hashlib
import io
import json
import re
import unicodedata
import uuid
from functools import lru_cache
from pathlib import Path
from string import Formatter

import arabic_reshaper
import reportlab
from bidi.algorithm import get_display
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_DIRECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.shared import Mm, Pt
from lxml import etree
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas

ASSETS = Path(__file__).resolve().parents[1] / "assets" / "m53"
SPEC = json.loads((ASSETS / "source_format.json").read_text(encoding="utf-8"))
CATALOGUE = json.loads((ASSETS / "variable-catalogue.json").read_text(encoding="utf-8"))
LABELS = {field["key"]: field["label_ar"] for field in CATALOGUE}
FIELDS = set(LABELS)
AUTOMATIC = {"rendered_page_count_words"}
REVISION = "m53-1"
ROWS, ROW_PT, MAX_PAGES = 28, 23.25, 100
MAIN_WIDTH = 129 * mm
REQUIRED = {"opening_time_words", "weekday", "day_words", "month_words", "year_words", "year",
            "subject", "opening_location", "investigator_name", "investigator_rank",
            "investigator_unit", "statement_name", "statement_signature_role", "duty_unit"}
for _key in ("investigator", "opening_record", "statement_heading", "identity", "signature_place"):
    REQUIRED.update(name for _, name, _, _ in Formatter().parse(SPEC["clauses"][_key]) if name)
CONFIRMATIONS = {
    "content_reviewed": "راجعت بيانات المحضر ونصوص الأسئلة والأجوبة ونسبتها إلى أصحابها",
    "clerk_appointment_confirmed": "تم تعيين الكاتب العدلي المذكور",
    "oath_administered_confirmed": "تم تحليف اليمين المذكورة في النموذج فعلياً",
    "military_attire_confirmed": "تم التحقق من عبارة ارتداء اللباس العسكري",
    "identity_check_confirmed": "تم التحقق من وثائق الهوية المذكورة وإعادتها",
    "read_back_confirmed": "تليت الإفادة على صاحبها",
    "attestation_confirmed": "صدّق صاحب الإفادة على مضمونها",
    "signature_recorded_confirmed": "تم تسجيل التوقيع فعلياً خارج هذا النموذج",
    "procedural_note_confirmed": "تمت المراجعة والقرار المذكوران في الملاحظة",
    "interim_closing_confirmed": "تم تنظيم وختم التحقيق الأولي كما ورد",
    "human_conclusion_confirmed": "راجعت الاستنتاج الذي أدخلته بنفسي",
    "human_responsibilities_confirmed": "راجعت المسؤوليات التي أدخلتها بنفسي",
    "human_proposals_confirmed": "راجعت الاقتراحات التي أدخلتها بنفسي",
    "final_closing_confirmed": "تم تنظيم وختم الملف كما ورد",
}
GROUPS = [
    ("افتتاح المحضر", ["opening_datetime", "investigator", "clerk", "opening_record"]),
    ("صاحب الإفادة", ["statement_heading", "identity", "identity_check"]),
    ("المراجعة والإقفال", ["signature_place", "note", "interim_closing"]),
    ("الاستنتاج والمسؤوليات والاقتراحات", ["conclusion", "responsibilities", "proposals"]),
]


def digest(value) -> str:
    return hashlib.sha256(json.dumps(value, ensure_ascii=False, sort_keys=True,
                                     default=str).encode("utf-8")).hexdigest()


@lru_cache(maxsize=1)
def format_digest() -> str:
    return digest({"revision": REVISION, "spec": SPEC,
                   "renderer": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                   "fonts": [hashlib.sha256(p.read_bytes()).hexdigest()
                             for p in sorted(ASSETS.glob("*.ttf"))]})


def catalogue() -> dict:
    used = set()
    groups = []
    for title, clauses in GROUPS:
        names = []
        for clause in clauses:
            for _, key, _, _ in Formatter().parse(SPEC["clauses"][clause]):
                if key and key not in used and key not in AUTOMATIC:
                    names.append(key)
                    used.add(key)
        groups.append({"title": title, "fields": names})
    groups.insert(0, {"title": "رأس النموذج", "fields": sorted(FIELDS - used - AUTOMATIC)})
    return {"revision": REVISION, "fields": [dict(field, required=field["key"] in REQUIRED,
                automatic=field["key"] in AUTOMATIC) for field in CATALOGUE],
            "groups": groups, "clauses": SPEC["clauses"], "fixed_labels": SPEC["fixed_labels"],
            "confirmations": CONFIRMATIONS, "requirements": SPEC["confirmation_requirements"]}


def review_token(draft, context: dict) -> str:
    # Excludes request time and server-generated approval metadata. Includes the
    # identity/attribution currently resolved by the server, not just raw speaker IDs.
    return digest({"format": format_digest(), "form": (draft.official_form or {}).get("values", {}),
        "enabled": (draft.official_form or {}).get("enabled", False),
        "context": {key: value for key, value in context.items()
                    if key not in {"generated_at", "generated_by_name", "report_version", "template_version", "template_notice"}},
        "pins": draft.pinned_transcripts, "sources": [{"id": str(b.id),
            "question": b.question_source_text, "answer": b.answer_source_text,
            "segment_ids": b.source_segment_ids, "recordings": b.source_recording_ids,
            "start": b.start_seconds, "end": b.end_seconds} for b in draft.qa_blocks]})


def confirmed(draft, token: str) -> dict[str, bool]:
    events = (draft.official_form or {}).get("confirmations", {})
    return {key: bool(events.get(key, {}).get("content_sha256") == token) for key in CONFIRMATIONS}


def values_for(draft) -> dict[str, str]:
    return {key: (draft.official_form or {}).get("values", {}).get(key, "") for key in FIELDS}


def missing_fields(values: dict, flags: dict) -> list[str]:
    needed = set(REQUIRED)
    # A confirmed clause must have every variable filled. Unconfirmed clauses do
    # not become false assertions simply to make the historical form look full.
    for clause, requirements in SPEC["confirmation_requirements"].items():
        if all(flags.get(flag) for flag in requirements):
            needed.update(key for _, key, _, _ in Formatter().parse(SPEC["clauses"][clause]) if key)
    return [LABELS[key] for key in sorted(needed - AUTOMATIC) if not values.get(key, "").strip()]


def suggestions(context: dict) -> dict:
    result = {"subject": context["case_subject"] or context["session_title"],
              "opening_location": context["investigation_location"],
              "investigator_name": context["investigator_name"],
              "investigator_rank": context["investigator_rank"],
              "investigator_unit": context["investigator_unit"]}
    # Never choose the first of several people and call all answers theirs.
    answer_names = {b["answer_speaker"] for b in context["qa_blocks"] if b["answer"]}
    if len(answer_names) == 1:
        result["statement_name"] = next(iter(answer_names))
    date = context["report_date"] or context["session_date"]
    if date:
        result["year"] = date[:4]
    return {key: value for key, value in result.items() if value}


_reshaper = arabic_reshaper.ArabicReshaper(configuration={
    "delete_harakat": False, "shift_harakat_position": True, "support_ligatures": False})


@lru_cache(maxsize=1)
def register_fonts():
    for style in ("Regular", "Bold"):
        pdfmetrics.registerFont(TTFont("M53" + style, str(ASSETS / f"NotoNaskhArabic-{style}.ttf")))
    font_dir = Path(reportlab.__file__).parent / "fonts"
    for style, name in (("Regular", "Vera.ttf"), ("Bold", "VeraBd.ttf")):
        pdfmetrics.registerFont(TTFont("M53Latin" + style, str(font_dir / name)))


def visual(text: str) -> str:
    try:
        return get_display(_reshaper.reshape(text), base_dir="R")
    except (AssertionError, KeyError) as exc:
        raise ValueError("report_m53_unsupported_character") from exc


def width(text: str, bold=False, size=12) -> float:
    register_fonts()
    return sum(pdfmetrics.stringWidth(part, font, size) for part, font in font_runs(visual(text), bold))


def font_runs(text, bold=False):
    primary = "M53Bold" if bold else "M53Regular"
    fallback = "M53LatinBold" if bold else "M53LatinRegular"
    runs = []
    for char in text:
        font = primary if pdfmetrics.getFont(primary).face.charToGlyph.get(ord(char)) else fallback
        if not pdfmetrics.getFont(font).face.charToGlyph.get(ord(char)):
            raise ValueError("report_m53_unsupported_character")
        if runs and runs[-1][1] == font:
            runs[-1] = (runs[-1][0] + char, font)
        else:
            runs.append((char, font))
    return runs


def wrap(text: str, bold=False, size=12, available=MAIN_WIDTH - 8):
    """Measure shaped glyphs; split oversized tokens without dropping characters."""
    lines = []
    for paragraph in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        line = ""
        for word in re.findall(r"\S+", paragraph):
            candidate = (line + " " + word).lstrip()
            if width(candidate, bold, size) <= available:
                line = candidate
                continue
            if line:
                lines.append(line)
                line = ""
            # Combining marks stay with their base character.
            clusters = []
            for char in word:
                if unicodedata.combining(char) and clusters:
                    clusters[-1] += char
                else:
                    clusters.append(char)
            for cluster in clusters:
                if line and width(line + cluster, bold, size) > available:
                    lines.append(line)
                    line = ""
                line += cluster
                if width(line, bold, size) > available:
                    raise ValueError("report_m53_text_too_wide")
        lines.append(line)
    return lines


def make_record(values, qa, flags, *, blank=False):
    return {"values": dict(values), "qa": qa, "flags": flags, "blank": blank}


def clause(record, key):
    if not record["blank"] and not all(record["flags"].get(flag)
            for flag in SPEC["confirmation_requirements"].get(key, [])):
        return ""
    return SPEC["clauses"][key].format_map(record["values"])


def paginate(record):
    labels = SPEC["fixed_labels"]
    def empty():
        return [{"text": "", "margin": "", "bold": False} for _ in range(ROWS)]
    pages = [empty()]
    pages[0][0]["text"] = labels["title"]
    for i, key in enumerate(("country", "ministry", "service")):
        pages[0][i]["margin"] = labels[key]
    pages[0][3]["margin"] = labels["duty"] + " " + record["values"]["duty_unit"]
    pages[0][4]["margin"] = labels["subject"]
    cursor = 2

    def new_page():
        nonlocal cursor
        if len(pages) >= MAX_PAGES:
            raise ValueError("report_too_large")
        pages.append(empty())
        pages[-1][0]["margin"] = labels["margin"]
        cursor = 1

    def write(text, margin="", bold=False):
        nonlocal cursor
        if not text:
            return
        for i, fragment in enumerate(wrap(text, bold)):
            if cursor >= ROWS:
                new_page()
            pages[-1][cursor].update(text=fragment, bold=bold)
            if i == 0 and margin:
                pages[-1][cursor]["margin"] = margin
            cursor += 1

    def anchor(page, row):
        nonlocal cursor
        while len(pages) - 1 < page:
            new_page()
        if len(pages) - 1 == page:
            cursor = max(cursor, row)

    def group(*texts):
        needed = sum(len(wrap(text)) for text in texts if text)
        if cursor + needed > ROWS and needed < ROWS:
            new_page()
        for text in texts:
            write(text)

    def section(key, label):
        # Keep the official section label even when its factual wording is not
        # confirmed. The empty ruled line remains available for later completion.
        write(clause(record, key) or " ", label)

    for key in ("opening_datetime", "investigator", "clerk", "oath", "military_attire", "opening_record"):
        write(clause(record, key))
    write(clause(record, "statement_heading"), bold=True)
    write(clause(record, "identity"), labels["statement_marker"])
    write(clause(record, "identity_check"))
    for qa in record["qa"]:
        if qa.get("question"):
            write(qa["question"], "س")
        if qa.get("answer"):
            write(qa["answer"], "ج")
    anchor(1, 6)
    group(clause(record, "attestation"), clause(record, "signature_place"), clause(record, "investigator_signature"))
    anchor(1, 10)
    section("note", labels["note"])
    if clause(record, "note"):
        group(clause(record, "signature_place"), clause(record, "investigator_signature"))
    anchor(1, 16)
    write(clause(record, "interim_closing"))
    if clause(record, "interim_closing"):
        group(clause(record, "signature_place"), clause(record, "investigator_signature"))
    anchor(1, 21)
    section("conclusion", labels["conclusion"])
    anchor(1, 26)
    section("responsibilities", labels["responsibilities"])
    anchor(2, 4)
    section("proposals", labels["proposals"])
    group(clause(record, "final_closing"),
          record["values"]["closing_location"] + " في " + record["values"]["closing_date"],
          clause(record, "investigator_signature"))
    return pages


def layout(record):
    pages = paginate(record)
    for _ in range(6):
        record["values"]["rendered_page_count_words"] = str(len(pages)) + " صفحة"
        updated = paginate(record)
        if len(updated) == len(pages):
            return updated
        pages = updated
    raise ValueError("report_m53_pagination_failed")


def _xml(tag, **attrs):
    node = OxmlElement("w:" + tag)
    for key, value in attrs.items():
        node.set(qn("w:" + key), str(value))
    return node


def _paragraph(p, text, *, bold=False, size=12, center=False):
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT
    p._p.get_or_add_pPr().append(_xml("bidi"))
    if not center:
        p._p.get_or_add_pPr().find(qn("w:jc")).set(qn("w:val"), "start")
    p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = Pt(ROW_PT)
    run = p.add_run(text)
    run.font.name = "Arial"
    run.font.size = Pt(size)
    run.bold = bold
    props = run._r.get_or_add_rPr()
    props.get_or_add_rFonts().set(qn("w:cs"), "Noto Naskh Arabic")
    props.append(_xml("szCs", val=int(size * 2)))
    props.append(_xml("bCs", val=int(bold)))
    props.append(_xml("rtl"))


def _embed_word_fonts(doc):
    """Embed licensed full fonts using the OOXML 32-byte font obfuscation."""
    table = doc.part.part_related_by(RT.FONT_TABLE)
    root = etree.fromstring(table.blob)
    for family, folder, files in (
        ("Noto Naskh Arabic", ASSETS, ("NotoNaskhArabic-Regular.ttf", "NotoNaskhArabic-Bold.ttf")),
    ):
        font = _xml("font", name=family)
        root.append(font)
        for style, filename in zip(("embedRegular", "embedBold"), files):
            font_key = uuid.uuid5(uuid.NAMESPACE_URL, "military-stt/m53/" + filename)
            font_bytes = bytearray((folder / filename).read_bytes())
            key_bytes = font_key.bytes[::-1]
            for i in range(32):
                font_bytes[i] ^= key_bytes[i % 16]
            part = Part(PackURI("/word/fonts/" + filename + ".odttf"),
                        "application/vnd.openxmlformats-officedocument.obfuscatedFont",
                        bytes(font_bytes), doc.part.package)
            relation = table.relate_to(part, RT.FONT)
            embedded = _xml(style, fontKey="{" + str(font_key).upper() + "}", subsetted="false")
            embedded.set(qn("r:id"), relation)
            font.append(embedded)
    table._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    doc.settings.element.append(_xml("embedTrueTypeFonts"))


def render_docx(record, pages, *, preview=False):
    doc = Document()
    doc.core_properties.title = "محضر تحقيق عدلي — م ٥٣"
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(6), Mm(38)
    sec.top_margin, sec.bottom_margin = Mm(28), Mm(22)
    sec.header_distance, sec.footer_distance = Mm(5), Mm(12)
    sec.different_first_page_header_footer = True

    def plain_table(parent, widths):
        table = parent.add_table(rows=1, cols=len(widths), width=Mm(sum(widths)))
        table.autofit = False
        table.table_direction = WD_TABLE_DIRECTION.RTL
        for column, cell, size in zip(table.columns, table.rows[0].cells, widths):
            column.width = cell.width = Mm(size)
        return table

    def small_separator(p):
        p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(1)

    def page_number(p):
        _paragraph(p, "ص - ", center=True, size=10)
        p._p.append(_xml("fldSimple", instr="PAGE"))
        _paragraph(p, " -", center=True, size=10)

    for header, first in ((sec.first_page_header, True), (sec.header, False)):
        _paragraph(header.paragraphs[0], "مسودة للمراجعة" if preview else "", size=7, center=True)
        header.paragraphs[0].paragraph_format.line_spacing = Pt(9)
        head = plain_table(header, (37, 129))
        _paragraph(head.cell(0, 0).paragraphs[0], record["values"]["year"], size=11, center=True)
        _paragraph(head.cell(0, 1).paragraphs[0], record["values"]["subject"], size=11, center=True)
        borders = _xml("tblBorders")
        for side in ("bottom", "insideV"):
            borders.append(_xml(side, val="single", sz=6, color="000000"))
        head._tbl.tblPr.append(borders)
        small_separator(header.add_paragraph())
        nums = plain_table(header, (45, 76, 45))
        page_number(nums.cell(0, 1).paragraphs[0])
        _paragraph(nums.cell(0, 2).paragraphs[0], SPEC["fixed_labels"]["model"] if first else "", size=10)
    for footer, first in ((sec.first_page_footer, True), (sec.footer, False)):
        small_separator(footer.paragraphs[0])
        if first:
            roles = plain_table(footer, (55, 56, 55))
            for cell, role in zip(roles.rows[0].cells, (record["values"]["statement_signature_role"], "كاتب عدلي", record["values"]["investigator_rank"] + " محقق")):
                _paragraph(cell.paragraphs[0], role, size=10, center=True)
        page_number(footer.add_paragraph())
    for page_index, rows in enumerate(pages):
        if page_index:
            p = doc.add_paragraph()
            p.paragraph_format.page_break_before = True
            p.paragraph_format.line_spacing = Pt(1)
            p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(0)
        table = doc.add_table(rows=ROWS, cols=2)
        table.autofit = False
        table.table_direction = WD_TABLE_DIRECTION.RTL
        for col, size in zip(table.columns, (37, 129)):
            col.width = Mm(size)
        borders = _xml("tblBorders")
        for side in ("top", "bottom", "left", "right", "insideV", "insideH"):
            borders.append(_xml(side, val="single" if side == "insideH" else "double", sz=4, color="000000"))
        table._tbl.tblPr.append(borders)
        for i, row_data in enumerate(rows):
            row = table.rows[i]
            row.height, row.height_rule = Pt(ROW_PT), WD_ROW_HEIGHT_RULE.EXACTLY
            row._tr.get_or_add_trPr().append(_xml("cantSplit"))
            for cell, size, text in zip(row.cells, (37, 129), (row_data["margin"], row_data["text"])):
                cell.width = Mm(size)
                mar = _xml("tcMar")
                for side in ("top", "bottom", "start", "end"):
                    mar.append(_xml(side, w=0 if side in ("top", "bottom") else 35, type="dxa"))
                cell._tc.get_or_add_tcPr().append(mar)
                _paragraph(cell.paragraphs[0], text, bold=row_data["bold"] if size == 129 else True, size=12 if size == 129 else 10)
        if page_index == 0:
            cell = table.cell(0, 1).merge(table.cell(1, 1))
            cell.text = ""
            _paragraph(cell.paragraphs[0], SPEC["fixed_labels"]["title"], size=23, bold=True, center=True)
    out = io.BytesIO()
    _embed_word_fonts(doc)
    doc.save(out)
    return out.getvalue()


def render_pdf(record, pages, *, preview=False):
    register_fonts()
    out = io.BytesIO()
    canvas = Canvas(out, pagesize=A4, pageCompression=1)
    canvas.setTitle("محضر تحقيق عدلي — م ٥٣")
    canvas.setAuthor("Military STT")
    left, divider, right, top = 6 * mm, 135 * mm, 172 * mm, A4[1] - 28 * mm

    def text(value, x1, x2, baseline, *, bold=False, size=12, center=False):
        shaped = visual(value)
        runs = font_runs(shaped, bold)
        def measured():
            return sum(pdfmetrics.stringWidth(part, font, size) for part, font in runs)
        # The body has already been wrapped; bounded header/margin fields scale
        # only to 8pt, then fail visibly instead of cutting off any text.
        while measured() > x2 - x1 - 6 and size > 8:
            size -= .25
        if measured() > x2 - x1 - 6:
            raise ValueError("report_m53_text_too_wide")
        x = (x1 + x2 - measured()) / 2 if center else x2 - 3 - measured()
        for part, font in runs:
            canvas.setFont(font, size)
            canvas.drawString(x, baseline, part)
            x += pdfmetrics.stringWidth(part, font, size)

    for index, rows in enumerate(pages):
        if preview:
            text("مسودة للمراجعة — لا تتضمن توقيعاً رقمياً", left, right, A4[1] - 8 * mm, size=8, center=True)
        text(record["values"]["subject"], left, divider, A4[1] - 16 * mm, center=True, size=11)
        text(record["values"]["year"], divider, right, A4[1] - 16 * mm, center=True, size=11)
        canvas.setLineWidth(.6)
        canvas.line(left, A4[1] - 18 * mm, right, A4[1] - 18 * mm)
        canvas.line(divider, A4[1] - 10 * mm, divider, A4[1] - 18 * mm)
        text("ص - " + str(index + 1) + " -", left, right, A4[1] - 24 * mm, center=True, size=11)
        if index == 0:
            text(SPEC["fixed_labels"]["model"], left, left + 50 * mm, A4[1] - 24 * mm, size=10)
        bottom = top - ROWS * ROW_PT
        canvas.setLineWidth(.4)
        for row in range(ROWS + 1):
            # The merged two-row title has no horizontal rule through it.
            canvas.line(divider if index == 0 and row == 1 else left, top - row * ROW_PT, right, top - row * ROW_PT)
        for x in (left, divider, right):
            for offset in (-.8, .8):
                canvas.line(x + offset, bottom, x + offset, top)
        for y in (top, bottom):
            canvas.line(left, y + .8, right, y + .8)
        for i, row in enumerate(rows):
            baseline = top - i * ROW_PT - 16
            text(row["margin"], divider + 1, right - 1, baseline, bold=True, size=10, center=index == 0 and i < 3)
            if index == 0 and i == 0:
                text(row["text"], left, divider, top - 32, size=23, bold=True, center=True)
            else:
                text(row["text"], left, divider, baseline, bold=row["bold"])
        if index == 0:
            for i, role in enumerate((record["values"]["investigator_rank"] + " محقق", "كاتب عدلي", record["values"]["statement_signature_role"])):
                text(role, left + i * (right-left)/3, left + (i+1)*(right-left)/3, bottom - 16, size=10, center=True)
        text("- " + str(index + 1) + " -", left, right, 12 * mm, size=10, center=True)
        canvas.showPage()
    canvas.save()
    return out.getvalue()


def render_documents(values, qa, flags, *, preview=False, blank=False):
    if sum(len(item.get("question", "")) + len(item.get("answer", "")) for item in qa) > 200000:
        raise ValueError("report_too_large")
    record = make_record(values, qa, flags, blank=blank)
    pages = layout(record)
    return render_docx(record, pages, preview=preview), render_pdf(record, pages, preview=preview), {
        "revision": REVISION, "values": record["values"], "pages": pages,
        "omitted_clauses": [key for key in SPEC["confirmation_requirements"] if not clause(record, key)]}

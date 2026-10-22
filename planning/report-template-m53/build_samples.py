"""Build review-only model M53 documents from fictional JSON, without app/DB access.

Uses the application's existing python-docx dependency. The result is a reconstruction
for visual review, not an approved template or an integration with the live composer.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT, WD_TABLE_DIRECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor


FONT = "Arial"
LATIN_FONT = "Arial"
WIDTHS = (34, 146)  # first column is displayed at the right by bidiVisual
NOTICE = "نسخة للمراجعة — بيانات افتراضية — غير معتمدة"
BLANK_NOTICE = "قالب مُعاد البناء للمراجعة — غير معتمد"
BLANK = "........................................"
ROOT = Path(__file__).resolve().parent


def element(tag, **attributes):
    node = OxmlElement(tag)
    for name, value in attributes.items():
        node.set(qn(f"w:{name}"), str(value))
    return node


def arabic_number(value):
    return str(value).translate(str.maketrans("0123456789", "٠١٢٣٤٥٦٧٨٩"))


def format_run(run, *, size=12, bold=False, latin=False, color=None):
    run.font.name = LATIN_FONT if latin else FONT
    run.font.size = Pt(size if not latin else size * 0.86)
    run.bold = bold
    rpr = run._r.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for name in ("ascii", "hAnsi"):
        fonts.set(qn(f"w:{name}"), LATIN_FONT)
    fonts.set(qn("w:cs"), FONT)
    rpr.append(element("w:rtl", val="0" if latin else "1"))
    rpr.append(element("w:szCs", val=int(size * 2)))
    rpr.append(element("w:bCs", val="1" if bold else "0"))
    rpr.append(element("w:lang", val="ar-LB", bidi="ar-LB"))
    if color:
        run.font.color.rgb = RGBColor.from_string(color)
    return run


def paragraph(p, text="", *, size=12, bold=False, center=False, keep=False, color=None):
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.12
    p.paragraph_format.keep_with_next = keep
    # Keep short paragraphs together so continuation pages do not start with a
    # single trailing line. Very long paragraphs remain free to span pages.
    p.paragraph_format.keep_together = len(text) <= 700
    p.paragraph_format.widow_control = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.RIGHT
    p._p.get_or_add_pPr().append(element("w:bidi"))
    if not center:
        # Logical start is the right edge for bidi paragraphs. Physical "right"
        # is mirrored by this Word renderer inside a bidi table.
        p._p.get_or_add_pPr().find(qn("w:jc")).set(qn("w:val"), "start")
    # Keep Latin identifiers/time strings in LTR runs, without changing source text.
    for token in re.split(r"([A-Za-z0-9][A-Za-z0-9_:/.-]*)", text):
        if token:
            format_run(p.add_run(token), size=size, bold=bold,
                       latin=bool(re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9_:/.-]*", token)),
                       color=color)
    return p


def field(p, instruction):
    run = format_run(p.add_run(), size=9, latin=True)
    run._r.append(element("w:fldChar", fldCharType="begin"))
    command = element("w:instrText")
    command.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    command.text = f" {instruction} "
    run._r.append(command)
    run._r.append(element("w:fldChar", fldCharType="separate"))
    text = element("w:t")
    text.text = "1"
    run._r.append(text)
    run._r.append(element("w:fldChar", fldCharType="end"))


def table_layout(table, *, widths=WIDTHS, borders=True):
    table.autofit = False
    table.table_direction = WD_TABLE_DIRECTION.RTL
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for column, width in zip(table.columns, widths):
        column.width = Mm(width)
    pr = table._tbl.tblPr
    border = element("w:tblBorders")
    for name in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border.append(element(f"w:{name}", val="single" if borders else "nil",
                              sz="8" if name != "insideH" else "4",
                              color="222222" if name != "insideH" else "888888"))
    pr.append(border)
    margins = element("w:tblCellMar")
    for name, value in (("top", 65), ("bottom", 65), ("start", 110), ("end", 110)):
        margins.append(element(f"w:{name}", w=value, type="dxa"))
    pr.append(margins)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.width = Mm(width)
    return table


def no_split(row):
    row._tr.get_or_add_trPr().append(element("w:cantSplit"))


def body_row(table, label, content, *, speaker=None, keep=False):
    # Word ignores keepLines within a single very tall table row. Keep each
    # source paragraph in its own row so short paragraphs really stay together.
    # Attribution is printed once; subsequent rows remain part of that statement.
    if not keep and "\n" in content:
        first = None
        for index, text in enumerate(content.split("\n")):
            row = body_row(table, label if index == 0 else "", text,
                           speaker=speaker if index == 0 else None,
                           keep=len(text) <= 700)
            first = first or row
        return first
    row = table.add_row()
    for cell, width in zip(row.cells, WIDTHS):
        cell.width = Mm(width)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    paragraph(row.cells[0].paragraphs[0], label, bold=True, size=11)
    cell = row.cells[1]
    texts = content.split("\n") if content else [BLANK]
    if speaker:
        paragraph(cell.paragraphs[0], speaker, bold=True, size=10, keep=True)
        for text in texts:
            paragraph(cell.add_paragraph(), text)
    else:
        for index, text in enumerate(texts):
            paragraph(cell.paragraphs[0] if index == 0 else cell.add_paragraph(), text)
    if keep or len(content) <= 700:
        no_split(row)
    return row


def build_document(record, destination):
    doc = Document()
    doc.core_properties.title = "محضر تحقيق عدلي — نموذج م ٥٣ — نسخة للمراجعة"
    doc.core_properties.author = "M53 layout review"
    doc.core_properties.subject = "Fictional layout specimen; not an official issued report"
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(25), Mm(20)
    section.left_margin = section.right_margin = Mm(15)
    section.header_distance, section.footer_distance = Mm(9), Mm(9)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.paragraph_format.space_after = Pt(3)
    doc.settings.element.append(element("w:updateFields", val="true"))
    notice = BLANK_NOTICE if record["sample_kind"] == "blank" else NOTICE

    for header, continuation in ((section.first_page_header, False), (section.header, True)):
        paragraph(header.paragraphs[0], notice, size=9, center=True, color="8B302C")
        paragraph(header.add_paragraph(),
                  "الجيش نموذج م ٥٣" if not continuation else "محضر تحقيق عدلي — تابع — نموذج م ٥٣",
                  size=10, bold=True, center=True)
    for footer in (section.first_page_footer, section.footer):
        p = paragraph(footer.paragraphs[0], "صفحة ", size=9, center=True)
        field(p, "PAGE")
        format_run(p.add_run(" من "), size=9)
        field(p, "NUMPAGES")

    # A separate first-page title block allows the body to flow independently.
    heading = table_layout(doc.add_table(rows=1, cols=2))
    no_split(heading.rows[0])
    institution, title = heading.rows[0].cells
    for index, line in enumerate(record["institution_lines"]):
        paragraph(institution.paragraphs[0] if index == 0 else institution.add_paragraph(),
                  line, size=10, bold=True, center=True)
    title.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    paragraph(title.paragraphs[0], "محضر تحقيق عدلي", size=23, bold=True, center=True)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.space_before = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(2)
    format_run(spacer.add_run(""), size=2)

    body = table_layout(doc.add_table(rows=1, cols=2))
    top = body.rows[0]
    top._tr.get_or_add_trPr().append(element("w:tblHeader"))
    no_split(top)
    paragraph(top.cells[0].paragraphs[0], "هامش", size=10, bold=True, center=True)
    paragraph(top.cells[1].paragraphs[0], "", size=5)
    body_row(body, "موضوع القضية", record["subject"], keep=True)
    opening = (
        f"التاريخ: {record['session_date'] or BLANK}    الساعة: {record['opening_time'] or BLANK}\n"
        f"مكان التحقيق: {record['location'] or BLANK}\n"
        f"المحقق: {record['investigator'] or BLANK}\n"
        f"الكاتب: {record['clerk'] or BLANK}\n"
        f"الإحالة: {record['referral'] or BLANK}"
    )
    body_row(body, "افتتاح المحضر", opening, keep=True)
    for statement in record["statements"]:
        label = {"question": "س", "answer": "ج", "statement": "إفادة"}[statement["kind"]]
        body_row(body, label, statement["text"], speaker=statement["speaker"] or None)
    for label, key in (("ملاحظة", "notes"), ("الاستنتاج", "conclusion"),
                       ("المسؤوليات", "responsibilities"), ("الاقتراحات", "proposals"),
                       ("إقفال المحضر", "closing")):
        body_row(body, label, record[key])
    # No embedded signatures, automatic oath, read-back or attestation assertions.
    body_row(body, "التوقيعات", "صاحب الإفادة: ____________________\n"
             "الكاتب: ____________________\nالمحقق: ____________________", keep=True)
    doc.save(destination)
    verify_docx(destination, record)


def verify_docx(path, record):
    doc = Document(path)
    paragraphs = [p.text for table in doc.tables for row in table.rows
                  for cell in row.cells for p in cell.paragraphs]
    for item in record["statements"]:
        for original in item["text"].split("\n"):
            if original and paragraphs.count(original) != 1:
                raise AssertionError(f"Source paragraph not preserved exactly once: {item['id']}")
    for key in ("notes", "conclusion", "responsibilities", "proposals", "closing"):
        for original in record[key].split("\n"):
            if original and paragraphs.count(original) != 1:
                raise AssertionError(f"Human-entered paragraph not preserved exactly once: {key}")
    with ZipFile(path) as archive:
        xml = archive.read("word/document.xml").decode()
        assert "tblHeader" in xml and "bidiVisual" in xml
        assert not any(name.startswith("word/media/") for name in archive.namelist())
        assert "hRule=\"exact\"" not in xml
        assert "w:type=\"page\"" not in xml  # no forced page breaks
        for name in ("word/footer1.xml", "word/footer2.xml"):
            assert "NUMPAGES" in archive.read(name).decode()


def fixtures():
    blank = {
        "schema_version": 1, "sample_kind": "blank", "fictional": True,
        "template_status": "reconstructed_for_review",
        "institution_lines": ["الجمهورية اللبنانية", "وزارة الدفاع الوطني", "الشرطة العسكرية", "خفر: ................"],
        "subject": "", "session_date": "", "opening_time": "", "location": "",
        "investigator": "", "clerk": "", "referral": "",
        "statements": [{"id": "blank-statement", "kind": "statement", "speaker": "", "text": ""}],
        "notes": "", "conclusion": "", "responsibilities": "", "proposals": "", "closing": "",
    }
    short = {**blank, "sample_kind": "short", "subject": "اختبار تنسيق محضر — بيانات افتراضية",
             "session_date": "٢١ / ١٠ / ٢٠٢٦", "opening_time": "09:30",
             "location": "قاعة افتراضية للاختبار", "investigator": "المحقق التجريبي",
             "clerk": "الكاتب التجريبي", "referral": "",
             "statements": [
                 {"id": "short-question", "kind": "question", "speaker": "المحقق التجريبي",
                  "text": "يرجى وصف ترتيب المواد المستخدمة في الاجتماع التجريبي؟"},
                 {"id": "short-answer", "kind": "answer", "speaker": "المشارك التجريبي الأول",
                  "text": "وضعت الملفات التجريبية على الطاولة بحسب ألوانها، ثم راجعت ترتيبها مع المشاركين. تحمل الحافظة الرمز DEMO-42. هذه بيانات افتراضية أعدت لمراجعة تنسيق المستند ولا تصف واقعة حقيقية."},
                 {"id": "short-statement", "kind": "statement", "speaker": "المشارك التجريبي الثاني",
                  "text": "أضيف أن الأوراق المستخدمة خالية من البيانات الشخصية. وهذه إفادة مستقلة لا يسبقها سؤال ضمن هذا المثال."},
             ]}
    long = {**short, "sample_kind": "long", "subject": "اختبار محضر ممتد عبر صفحات — بيانات افتراضية"}
    long_statement = [
        f"المقطع التجريبي {arabic_number(i)}: في هذا المثال الافتراضي، رتبت المواد الورقية على طاولة العمل وراجعت عناوين الملفات مع المشارك التجريبي الثاني. ثم سجلت الملاحظات كما وردت دون إضافة نتائج أو تحديد مسؤوليات. والغرض من هذه الفقرة هو فحص استمرار الإفادة على الصفحة التالية مع بقاء النص كاملاً ومقروءاً."
        for i in range(1, 25)
    ]
    long["statements"] = [
        {"id": "long-question", "kind": "question", "speaker": "المحقق التجريبي",
         "text": "يرجى وصف خطوات ترتيب الملفات الافتراضية بالتفصيل؟"},
        {"id": "long-answer", "kind": "answer", "speaker": "المشارك التجريبي الأول",
         "text": "\n".join(long_statement)},
        {"id": "literal-statement", "kind": "statement", "speaker": "المشارك التجريبي الثاني",
         "text": "وردت على الورقة التجريبية العبارة الحرفية {{ sample }} والرمز DEMO-42 والقيم 09:30 و٢١/١٠/٢٠٢٦ وعلامتا < >. تبقى هذه الرموز نصاً كما وردت، دون أن تنفذ كتعليمات أو تستبدل بقيم أخرى."},
        {"id": "last-question", "kind": "question", "speaker": "المحقق التجريبي",
         "text": "هل توجد إضافة أخرى تخص المثال الافتراضي؟"},
        {"id": "last-answer", "kind": "answer", "speaker": "المشارك التجريبي الثاني",
         "text": "هذه هي الإضافة الأخيرة في المثال، وقد أدرجت لفحص ظهور نهاية الإفادات قبل الأقسام الختامية."},
    ]
    long["notes"] = "ملاحظة تجريبية كتبها مُعدّ المثال: تركت بيانات الإحالة والتصديق والتوقيع فارغة لعدم وجود إجراء حقيقي في هذا المستند."
    long["conclusion"] = "نص اصطناعي لمراجعة موضع الاستنتاج فقط. لا يتضمن هذا المثال نتيجة تحقيق أو حكماً على أي شخص."
    long["responsibilities"] = "\n".join(
        f"فقرة المسؤوليات التجريبية {arabic_number(i)}: هذا نص وضع يدوياً لفحص امتداد القسم وتطابق الهوامش عند الانتقال بين الصفحات. لا ينسب هذا المثال مسؤولية إلى شخص حقيقي، ولا يستخدم كمضمون لمحضر فعلي."
        for i in range(1, 7)
    )
    long["proposals"] = "مقترح تجريبي لمراجعة التصميم: التأكد من وضوح النص العربي واستمرار الأقسام دون حذف، ثم مراجعة مواضع الهوامش وأرقام الصفحات."
    return {"m53-blank-review": blank, "m53-short-example": short, "m53-long-example": long}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=ROOT / "output")
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=True)
    manifest = {"classification": "reconstructed_for_review", "fictional": True,
                "source": "Supplied PDF pages 18-20; no historical personal data copied",
                "font": FONT, "artifacts": [], "native_word_visual_review": "pending"}
    for name, record in fixtures().items():
        json_path = output / f"{name}.json"
        json_path.write_text(json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        docx = output / f"{name}.docx"
        build_document(record, docx)
        artifacts = [json_path, docx]
        for path in artifacts:
            manifest["artifacts"].append({"file": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
                                          "bytes": path.stat().st_size})
        print(f"Generated {name}", flush=True)
    (output / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    main()

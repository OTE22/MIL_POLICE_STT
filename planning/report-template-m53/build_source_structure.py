"""Source-faithful ruled M53 layout with fixed phrases and session variables.

This supersedes the earlier redesigned table specimen. Produces editable Word
review documents and a complete variable catalogue; no application/DB access.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from string import Formatter
from zipfile import ZipFile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.shared import Mm, Pt

import build_samples as shared

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "source-structure"
SPEC = json.loads((ROOT / "source_format.json").read_text(encoding="utf-8"))
LABELS = SPEC["fixed_labels"]
ROW_PT = 23.25
ROWS = 28
WIDTHS = (37, 129)
FIELDS = sorted({name for pattern in SPEC["clauses"].values()
                 for _, name, _, _ in Formatter().parse(pattern) if name} | {"year", "duty_unit"})
ARABIC_LABELS = {
    "opening_time_words": "الساعة", "weekday": "اليوم", "day_words": "اليوم كتابةً",
    "month_words": "الشهر", "year_words": "السنة كتابةً", "year": "السنة",
    "investigator_rank": "رتبة المحقق", "investigator_name": "اسم المحقق",
    "investigator_identifier": "معرّف المحقق", "investigator_unit": "وحدة المحقق",
    "clerk_rank": "رتبة الكاتب", "clerk_name": "اسم الكاتب", "clerk_service_number": "رقم الكاتب",
    "clerk_unit": "وحدة الكاتب", "opening_location": "مكان الافتتاح",
    "referral_number": "رقم الإحالة", "referral_date": "تاريخ الإحالة", "subject": "الموضوع",
    "statement_rank": "الرتبة", "statement_name": "الاسم", "statement_service_number": "الرقم",
    "father_name": "الأب", "mother_name": "الأم", "birth_year": "سنة الولادة",
    "birth_place": "مكان الولادة", "district": "القضاء", "civil_register": "السجل",
    "residence": "السكن", "phone": "الهاتف", "family_status": "الحالة العائلية",
    "children_statement": "الأولاد", "education": "التعليم", "occupation_statement": "العمل والخدمة",
    "prior_record_statement": "إفادة السوابق", "identity_documents": "وثائق التحقق",
    "statement_signature_role": "صفة المستجوب", "closing_location": "مكان الإقفال",
    "closing_date": "تاريخ الإقفال", "note_date": "تاريخ المراجعة", "note_time": "وقتها",
    "consulted_unit": "القيادة", "consulted_rank": "الرتبة", "consulted_name": "الاسم",
    "procedural_decision": "القرار المدخل يدوياً", "copy_count_words": "عدد النسخ كتابةً",
    "interim_page_count_words": "عدد صفحات التحقيق", "rendered_page_count_words": "عدد صفحات الملف",
    "attachments": "المرفقات", "human_conclusion": "الاستنتاج المدخل يدوياً",
    "responsible_rank": "الرتبة", "responsible_name": "الاسم", "responsible_service_number": "الرقم",
    "responsible_unit": "الوحدة", "human_responsibilities": "المسؤولية المدخلة يدوياً",
    "human_proposals": "الاقتراحات المدخلة يدوياً", "duty_unit": "الخفر"
}
assert set(ARABIC_LABELS) == set(FIELDS)


def wrap(text, limit=63):
    """Conservative editable line wrapping; PDF character checks catch any clipping."""
    lines, line = [], ""
    for token in text.split():
        if line and len(line) + 1 + len(token) > limit:
            lines.append(line)
            line = token
        else:
            line = (line + " " + token).strip()
    if line:
        lines.append(line)
    return lines or [""]


def line(p, text="", *, bold=False, center=False, size=12):
    shared.paragraph(p, text, size=size, bold=bold, center=center)
    pf = p.paragraph_format
    pf.space_before = pf.space_after = Pt(0)
    pf.line_spacing = Pt(ROW_PT)
    pf.keep_together = True
    pf.keep_with_next = False


def clause(record, key):
    if record["kind"] != "blank":
        requirements = SPEC["confirmation_requirements"].get(key, [])
        if not all(record["confirmations"].get(flag) is True for flag in requirements):
            return ""
    return SPEC["clauses"][key].format_map(record["values"])


def paginate(record):
    # Each page is the same 28-line ruled sheet; the first has a two-line title.
    pages = [[{"text": "", "margin": "", "bold": False} for _ in range(ROWS)]]
    pages[0][0]["text"] = LABELS["title"]
    for i, key in enumerate(("country", "ministry", "service")):
        pages[0][i]["margin"] = LABELS[key]
    pages[0][3]["margin"] = LABELS["duty"] + " " + record["values"]["duty_unit"]
    pages[0][4]["margin"] = LABELS["subject"]
    cursor = 2

    def new_page():
        nonlocal cursor
        pages.append([{"text": "", "margin": "", "bold": False} for _ in range(ROWS)])
        pages[-1][0]["margin"] = LABELS["margin"]
        cursor = 1

    def write(text, margin="", bold=False):
        nonlocal cursor
        if not text:
            return
        for index, fragment in enumerate(wrap(text)):
            if cursor >= ROWS:
                new_page()
            row = pages[-1][cursor]
            row["text"], row["bold"] = fragment, bold
            if index == 0 and margin:
                row["margin"] = margin
            cursor += 1

    def source_anchor(page_number, row_number):
        """Preserve the reference's baseline section positions, allowing overflow.

        Longer input may move a section later; it can never overwrite prior text.
        Empty lines remain ruled, as in the original paper form.
        """
        nonlocal cursor
        while len(pages) - 1 < page_number:
            new_page()
        if len(pages) - 1 == page_number and cursor < row_number:
            cursor = row_number

    def write_group(*texts):
        required = sum(len(wrap(text)) for text in texts if text)
        if cursor + required > ROWS and required < ROWS:
            new_page()
        for text in texts:
            write(text)

    for key in ("opening_datetime", "investigator", "clerk", "oath", "military_attire", "opening_record"):
        write(clause(record, key))
    write(clause(record, "statement_heading"), bold=True)
    write(clause(record, "identity"), LABELS["statement_marker"])
    write(clause(record, "identity_check"))
    for statement in record["statements"]:
        write(statement, LABELS["statement_marker"])
    # Preserve the source's sequence: statement/read-back and signatures, note
    # and signatures, interim closure, then findings/responsibilities/proposals.
    source_anchor(1, 6)
    write_group(clause(record, "attestation"), clause(record, "signature_place"), clause(record, "investigator_signature"))
    source_anchor(1, 10)
    write(clause(record, "note"), LABELS["note"])
    write_group(clause(record, "signature_place"), clause(record, "investigator_signature"))
    source_anchor(1, 16)
    write(clause(record, "interim_closing"))
    write_group(clause(record, "signature_place"), clause(record, "investigator_signature"))
    source_anchor(1, 21)
    write(clause(record, "conclusion"), LABELS["conclusion"])
    source_anchor(1, 26)
    write(clause(record, "responsibilities"), LABELS["responsibilities"])
    source_anchor(2, 4)
    write(clause(record, "proposals"), LABELS["proposals"])
    write(clause(record, "final_closing"))
    write_group(record["values"]["closing_location"] + " في " + record["values"]["closing_date"],
                clause(record, "investigator_signature"))
    return pages


def set_cell(cell, text, *, bold=False, center=False, size=12):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    line(cell.paragraphs[0], text, bold=bold, center=center, size=size)
    # Narrow symmetric margins match the ruled reference rather than form cards.
    props = cell._tc.get_or_add_tcPr()
    margins = shared.element("w:tcMar")
    for side, value in (("top", 0), ("bottom", 0), ("start", 35), ("end", 35)):
        margins.append(shared.element(f"w:{side}", w=value, type="dxa"))
    props.append(margins)


def build(record, target):
    pages = paginate(record)
    if record["kind"] != "blank":
        # Actual rendered page count is a variable; reflow to a stable result.
        for _ in range(5):
            count_words = {1: "صفحة واحدة", 2: "صفحتين", 3: "ثلاث صفحات", 4: "أربع صفحات",
                           5: "خمس صفحات", 6: "ست صفحات", 7: "سبع صفحات", 8: "ثماني صفحات",
                           9: "تسع صفحات", 10: "عشر صفحات"}
            record["values"]["rendered_page_count_words"] = count_words.get(len(pages), shared.arabic_number(len(pages)) + " صفحة")
            updated = paginate(record)
            if len(updated) == len(pages):
                pages = updated
                break
            pages = updated
        else:
            raise ValueError("Page count did not stabilize")
    doc = Document()
    doc.core_properties.title = "محضر تحقيق عدلي — قالب بصياغة المصدر — للمراجعة"
    doc.core_properties.author = "M53 template review"
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.left_margin, sec.right_margin = Mm(6), Mm(38)
    sec.top_margin, sec.bottom_margin = Mm(28), Mm(22)
    sec.header_distance, sec.footer_distance = Mm(5), Mm(18)
    sec.different_first_page_header_footer = True
    for header, first in ((sec.first_page_header, True), (sec.header, False)):
        p = header.paragraphs[0]
        shared.paragraph(p, "إعادة بناء للمراجعة — بيانات افتراضية / حقول فارغة — غير معتمد",
                         size=7, center=True, color="8B302C")
        p.paragraph_format.space_after = Pt(0)
        head = shared.table_layout(header.add_table(rows=1, cols=2, width=Mm(166)), widths=WIDTHS, borders=False)
        set_cell(head.cell(0, 0), record["values"]["year"], center=True, size=11)
        set_cell(head.cell(0, 1), record["values"]["subject"], center=True, size=11)
        border = shared.element("w:tblBorders")
        border.append(shared.element("w:bottom", val="single", sz=8, color="000000"))
        border.append(shared.element("w:insideV", val="single", sz=6, color="000000"))
        for old in list(head._tbl.tblPr.findall(shared.qn("w:tblBorders"))):
            head._tbl.tblPr.remove(old)
        head._tbl.tblPr.append(border)
        separator = header.add_paragraph()
        separator.paragraph_format.space_before = separator.paragraph_format.space_after = Pt(0)
        separator.paragraph_format.line_spacing = Pt(1)
        nums = shared.table_layout(header.add_table(rows=1, cols=3, width=Mm(166)), widths=(45, 76, 45), borders=False)
        set_cell(nums.cell(0, 0), "", size=10)
        p = nums.cell(0, 1).paragraphs[0]
        shared.paragraph(p, "ص - ", size=11, center=True)
        shared.field(p, "PAGE")
        shared.format_run(p.add_run(" -"), size=11)
        set_cell(nums.cell(0, 2), LABELS["model"] if first else "", size=10)
    for footer, first in ((sec.first_page_footer, True), (sec.footer, False)):
        if first:
            footer.paragraphs[0].paragraph_format.line_spacing = Pt(1)
            roles = shared.table_layout(footer.add_table(rows=1, cols=3, width=Mm(166)), widths=(55, 56, 55), borders=False)
            set_cell(roles.cell(0, 0), record["values"]["statement_signature_role"], size=10, center=True)
            set_cell(roles.cell(0, 1), "كاتب عدلي", size=10, center=True)
            set_cell(roles.cell(0, 2), record["values"]["investigator_rank"] + " محقق", size=10, center=True)
            p = footer.add_paragraph()
        else:
            p = footer.paragraphs[0]
        shared.paragraph(p, "- ", size=10, center=True)
        shared.field(p, "PAGE")
        shared.format_run(p.add_run(" -"), size=10)
    for index, page in enumerate(pages):
        if index:
            doc.add_page_break()
        table = shared.table_layout(doc.add_table(rows=ROWS, cols=2), widths=WIDTHS)
        # Override default grey form-row dividers with continuous black ruling.
        borders = shared.element("w:tblBorders")
        for tag in ("top", "bottom", "left", "right", "insideV", "insideH"):
            borders.append(shared.element(f"w:{tag}", val="double" if tag != "insideH" else "single", sz=6 if tag != "insideH" else 4, color="000000"))
        for old in list(table._tbl.tblPr.findall(shared.qn("w:tblBorders"))):
            table._tbl.tblPr.remove(old)
        table._tbl.tblPr.append(borders)
        for row_index, data in enumerate(page):
            row = table.rows[row_index]
            row.height, row.height_rule = Pt(ROW_PT), WD_ROW_HEIGHT_RULE.EXACTLY
            shared.no_split(row)
            set_cell(row.cells[0], data["margin"], bold=True, size=11,
                     center=row_index < 3 if index == 0 else row_index == 0)
            set_cell(row.cells[1], data["text"], bold=data["bold"], size=12)
        if index == 0:
            title = table.cell(0, 1).merge(table.cell(1, 1))
            for p in title.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            title.paragraphs[0].clear()
            set_cell(title, LABELS["title"], center=True, bold=True, size=24)
        # Keep the page-break paragraph small and avoid an extra blank page.
        if index < len(pages) - 1:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.line_spacing = Pt(1)
    doc.save(target)
    # Exact strings must survive in the Word XML. PDF checks subsequently detect
    # clipped/overflowing fixed-grid cells rather than allowing silent truncation.
    loaded = Document(target)
    found = [cell.text for table in loaded.tables for row in table.rows for cell in row.cells]
    for page in pages:
        for data in page:
            if data["text"] and data["text"] not in found:
                raise AssertionError("A placed source line is missing from the DOCX")
    with ZipFile(target) as z:
        assert not any(n.startswith("word/media/") for n in z.namelist())
    return pages


def records():
    values = {key: "[" + ARABIC_LABELS[key] + "]" for key in FIELDS}
    flags = sorted({flag for req in SPEC["confirmation_requirements"].values() for flag in req})
    blank = {"kind": "blank", "fictional": True, "values": values,
             "confirmations": {flag: False for flag in flags},
             "statements": ["[نص الإفادة المراجع كما ورد من صاحبها]", "[جواب مراجع]", "[جواب مراجع]", "[الإضافة الأخيرة إن وجدت]"]}
    sample_values = {key: "بيان افتراضي" for key in FIELDS}
    sample_values.update(opening_time_words="التاسعة والنصف", weekday="الأربعاء", day_words="الحادي والعشرين",
                         month_words="تشرين الأول", year_words="ألفين وست وعشرين", year="٢٠٢٦",
                         subject="واقعة افتراضية لاختبار القالب", investigator_rank="رتبة افتراضية",
                         investigator_name="المحقق التجريبي", investigator_identifier="تجريبي",
                         investigator_unit="وحدة افتراضية", clerk_name="الكاتب التجريبي",
                         clerk_service_number="تجريبي", clerk_unit="الوحدة الافتراضية",
                         opening_location="مكتب افتراضي", referral_number="تجريبي", referral_date="٢١/١٠/٢٠٢٦",
                         statement_name="المشارك التجريبي", statement_service_number="تجريبي", birth_year="١٩٩٠",
                         phone="غير مدخل", note_date="٢١/١٠/٢٠٢٦", note_time="١٢:٠٠",
                         closing_date="٢١/١٠/٢٠٢٦", closing_location="مكان افتراضي",
                         copy_count_words="نسخة واحدة", interim_page_count_words="[عدد صفحات هذا الجزء]",
                         human_conclusion="[استنتاج يكتبه المراجع]", human_responsibilities="[مسؤوليات يحددها المراجع]",
                         human_proposals="[اقتراحات يكتبها المراجع]", procedural_decision="[قرار إجرائي يؤكده المراجع]",
                         attachments="[المرفقات الفعلية]", duty_unit="افتراضي")
    sample_values.update(investigator_rank="نقيب", clerk_rank="رقيب", clerk_unit="الوحدة الافتراضية",
                         statement_rank="رقيب أول", statement_signature_role="رقيب أول مستجوب",
                         father_name="الأب التجريبي", mother_name="الأم التجريبية", birth_place="البلدة الافتراضية",
                         district="القضاء الافتراضي", civil_register="تجريبي", residence="عنوان افتراضي",
                         family_status="متأهل", children_statement="وله ولدان", education="تعليماً افتراضياً",
                         occupation_statement="أعمل بوظيفة افتراضية في وحدة تدريبية",
                         prior_record_statement="[إفادة السوابق يراجعها المحقق]",
                         identity_documents="[وثائق راجعها المحقق]", consulted_unit="الوحدة الافتراضية",
                         consulted_rank="رتبة افتراضية", consulted_name="المسؤول التجريبي",
                         responsible_rank="[الرتبة]", responsible_name="[الاسم]", responsible_service_number="[الرقم]",
                         responsible_unit="[الوحدة]")
    # All confirmations are explicitly fictional; no real report is approved.
    short = {"kind": "filled", "fictional": True, "values": sample_values,
             "confirmations": {flag: True for flag in flags},
             "statements": [
                 "أفيد أنني حضرت إلى المكتب الافتراضي في الموعد المحدد، وقد رتبت المواد التجريبية وراجعت عناوينها. هذه بيانات افتراضية لا تصف واقعة حقيقية.",
                 "نعم قمت بمراجعة ترتيب الملفات التجريبية ولم أضف أي معلومات شخصية إليها.",
                 "كلا ليس لدي ما أضيفه على إفادتي التجريبية هذه."
             ]}
    long = {**short, "kind": "long", "values": dict(sample_values), "statements": [
        f"المقطع الافتراضي {shared.arabic_number(i)}: أفيد أنني راجعت المواد التجريبية على الطاولة، ثم سجلت الملاحظات كما وردت دون إضافة نتائج أو تحديد مسؤوليات. هذا نص افتراضي لاختبار استمرار الإفادة مع المحافظة على صياغة النموذج وهوامشه."
        for i in range(1, 15)
    ]}
    return {"m53-source-blank": blank, "m53-source-filled": short, "m53-source-long": long}


def main():
    OUT.mkdir(exist_ok=True)
    catalogue = [{"key": key, "label_ar": ARABIC_LABELS[key],
                  "kind": "derived_page_count" if key == "rendered_page_count_words" else "reviewed_variable"}
                 for key in FIELDS]
    (OUT / "variable-catalogue.json").write_text(json.dumps(catalogue, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    for name, record in records().items():
        pages = build(record, OUT / f"{name}.docx")
        (OUT / f"{name}.json").write_text(json.dumps(record, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        (OUT / f"{name}.layout.json").write_text(json.dumps(pages, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        print(f"Generated {name}: {len(pages)} ruled pages, {len(FIELDS)} variables", flush=True)
    # A directly readable writing blueprint, with no historical case values.
    text = "# الصياغة الثابتة والحقول المتغيرة — نموذج م ٥٣\n\nإعادة بناء للمراجعة من الصفحات ١٨–٢٠؛ ليست شهادة اعتماد.\n\n"
    for key, pattern in SPEC["clauses"].items():
        text += f"## {key}\n\n{pattern}\n\n"
        if key in SPEC["confirmation_requirements"]:
            text += "لا يدرج في محضر مكتمل إلا بعد تأكيد الوقائع المرتبطة به.\n\n"
    (OUT / "generic-writing-format.md").write_text(text, encoding="utf-8")


if __name__ == "__main__":
    main()
